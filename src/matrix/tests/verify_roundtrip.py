# Orijinal ve kaydedilmiş xlsx'i openpyxl ile karşılaştırır.  python3 verify_roundtrip.py orig.xlsx saved.xlsx
import sys
from openpyxl import load_workbook

a = load_workbook(sys.argv[1])
b = load_workbook(sys.argv[2])
bv = load_workbook(sys.argv[2], data_only=True)
problems = []
def chk(cond, msg):
    if not cond: problems.append(msg)

chk(a.sheetnames == b.sheetnames, f'sayfa adları: {a.sheetnames} != {b.sheetnames}')
for name in a.sheetnames:
    if name not in b.sheetnames: continue
    wa, wb_, wv = a[name], b[name], bv[name]
    chk(wa.sheet_state == wb_.sheet_state, f'{name}: durum {wa.sheet_state} != {wb_.sheet_state}')
    ta = wa.sheet_properties.tabColor.rgb if wa.sheet_properties.tabColor else None
    tb = wb_.sheet_properties.tabColor.rgb if wb_.sheet_properties.tabColor else None
    chk((ta or '')[-6:] == (tb or '')[-6:], f'{name}: sekme rengi {ta} != {tb}')
    chk(wa.freeze_panes == wb_.freeze_panes, f'{name}: donmuş bölme {wa.freeze_panes} != {wb_.freeze_panes}')
    ma = sorted(str(m) for m in wa.merged_cells.ranges); mb = sorted(str(m) for m in wb_.merged_cells.ranges)
    chk(ma == mb, f'{name}: birleştirmeler {ma} != {mb}')
    for col, d in wa.column_dimensions.items():
        if d.width and d.customWidth:
            wb_w = wb_.column_dimensions[col].width
            chk(wb_w and abs(wb_w - d.width) < 0.6, f'{name}: sütun {col} genişlik {d.width} != {wb_w}')
        if d.hidden: chk(wb_.column_dimensions[col].hidden, f'{name}: sütun {col} gizli değil')
    for r, d in wa.row_dimensions.items():
        if d.height: chk(wb_.row_dimensions[r].height and abs(wb_.row_dimensions[r].height - d.height) < 1, f'{name}: satır {r} yükseklik {d.height} != {wb_.row_dimensions[r].height}')
        if d.hidden: chk(wb_.row_dimensions[r].hidden, f'{name}: satır {r} gizli değil')
    for row in wa.iter_rows():
        for c in row:
            if c.value is None: continue
            cb = wb_[c.coordinate]
            va, vb = c.value, cb.value
            if isinstance(va, str) and va.startswith('='):
                chk(isinstance(vb, str) and vb.replace(' ', '').upper() == va.replace(' ', '').upper(), f'{name}!{c.coordinate}: formül {va!r} != {vb!r}')
                chk(wv[c.coordinate].value is not None, f'{name}!{c.coordinate}: formül sonucu (önbellek) yok')
            else:
                if hasattr(va, 'year') and hasattr(vb, 'year'):
                    chk(abs((va - vb).total_seconds()) < 1 if hasattr(va, 'hour') and hasattr(vb, 'hour') else str(va)[:10] == str(vb)[:10], f'{name}!{c.coordinate}: tarih {va!r} != {vb!r}')
                elif isinstance(va, float) or isinstance(vb, float):
                    chk(isinstance(vb, (int, float)) and abs(va - vb) < 1e-9, f'{name}!{c.coordinate}: değer {va!r} != {vb!r}')
                else:
                    chk(va == vb, f'{name}!{c.coordinate}: değer {va!r} != {vb!r}')
            if c.number_format != 'General':
                chk(c.number_format == cb.number_format, f'{name}!{c.coordinate}: sayı biçimi {c.number_format!r} != {cb.number_format!r}')
            fa, fb = c.font, cb.font
            chk(bool(fa.b) == bool(fb.b), f'{name}!{c.coordinate}: kalın {fa.b} != {fb.b}')
            chk(bool(fa.i) == bool(fb.i), f'{name}!{c.coordinate}: italik')
            chk(bool(fa.strike) == bool(fb.strike), f'{name}!{c.coordinate}: üstü çizili')
            chk((fa.u or None) == (fb.u or None), f'{name}!{c.coordinate}: altı çizili {fa.u} != {fb.u}')
            chk((fa.name or 'Calibri') == (fb.name or 'Calibri'), f'{name}!{c.coordinate}: yazı tipi {fa.name} != {fb.name}')
            chk(float(fa.sz or 11) == float(fb.sz or 11), f'{name}!{c.coordinate}: boyut {fa.sz} != {fb.sz}')
            ca = fa.color.rgb if fa.color is not None and fa.color.type == 'rgb' else None
            cbb = fb.color.rgb if fb.color is not None and fb.color.type == 'rgb' else None
            if ca and ca[-6:] != '000000': chk(cbb and cbb[-6:] == ca[-6:], f'{name}!{c.coordinate}: yazı rengi {ca} != {cbb}')
            if c.fill and c.fill.fill_type == 'solid' and c.fill.fgColor.type == 'rgb':
                chk(cb.fill.fill_type == 'solid' and cb.fill.fgColor.rgb[-6:] == c.fill.fgColor.rgb[-6:], f'{name}!{c.coordinate}: dolgu {c.fill.fgColor.rgb} != {cb.fill.fgColor.rgb}')
            for side in ('left', 'right', 'top', 'bottom'):
                sa, sb = getattr(c.border, side), getattr(cb.border, side)
                if sa is not None and sa.style:
                    chk(sb is not None and sb.style == sa.style, f'{name}!{c.coordinate}: kenarlık {side} {sa.style} != {sb.style if sb else None}')
            al, bl = c.alignment, cb.alignment
            chk((al.horizontal or None) == (bl.horizontal or None), f'{name}!{c.coordinate}: yatay hizalama {al.horizontal} != {bl.horizontal}')
            chk((al.vertical or None) in ((bl.vertical or None), None), f'{name}!{c.coordinate}: dikey hizalama {al.vertical} != {bl.vertical}')
            chk(bool(al.wrap_text) == bool(bl.wrap_text), f'{name}!{c.coordinate}: metin kaydırma')
            chk(int(al.text_rotation or 0) == int(bl.text_rotation or 0), f'{name}!{c.coordinate}: döndürme {al.text_rotation} != {bl.text_rotation}')
            if c.hyperlink: chk(cb.hyperlink is not None and cb.hyperlink.target == c.hyperlink.target, f'{name}!{c.coordinate}: bağlantı {c.hyperlink.target} != {cb.hyperlink.target if cb.hyperlink else None}')
            if c.comment: chk(cb.comment is not None and cb.comment.text.strip() == c.comment.text.strip(), f'{name}!{c.coordinate}: not {c.comment.text!r} != {cb.comment.text if cb.comment else None!r}')
    dva = [(str(d.sqref), d.type, d.formula1) for d in wa.data_validations.dataValidation]
    dvb = [(str(d.sqref), d.type, d.formula1) for d in wb_.data_validations.dataValidation]
    chk(sorted(dva) == sorted(dvb), f'{name}: veri doğrulama {dva} != {dvb}')
na = sorted((n, d.attr_text) for n, d in a.defined_names.items())
nb = sorted((n, d.attr_text) for n, d in b.defined_names.items())
chk(na == nb, f'tanımlı adlar {na} != {nb}')
chk(a.active.title == b.active.title, f'etkin sayfa {a.active.title} != {b.active.title}')
print('SORUN YOK' if not problems else f'{len(problems)} sorun:')
for p in problems: print(' -', p)
# formül sonuçları örneği
for name in b.sheetnames[:2]:
    ws = bv[name]
    print(name, [(c.coordinate, c.value) for row in ws.iter_rows(max_row=12, max_col=6) for c in row if c.value is not None and wb_ is not None][:40])
