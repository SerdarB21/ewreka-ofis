# openpyxl çıktısını Excel'in yazdığına benzer hale getirir: paylaşılan formüller, tema renkleri, önbellek değerleri
import re, sys, zipfile, shutil
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.styles.colors import Color
src = '/tmp/claude-0/excel_like_base.xlsx'
out = sys.argv[1] if len(sys.argv) > 1 else '/home/claude/ewreka-ofis/test-files/excel-benzeri.xlsx'
wb = Workbook(); ws = wb.active; ws.title = 'Veri'
ws['A1'] = 'Değer'; ws['B1'] = 'İki katı'; ws['C1'] = 'Tema rengi'
for r in range(2, 12):
    ws.cell(row=r, column=1, value=r * 3)
    ws.cell(row=r, column=2, value=f'=A{r}*2')
ws['C2'] = 'Vurgu 1 açık'
ws['C2'].font = Font(color=Color(theme=4, tint=0.3999), bold=True)
ws['C3'] = 'Dolgu tema 9'
ws['C3'].fill = PatternFill('solid', fgColor=Color(theme=9, tint=-0.249))
ws['D2'] = '=XLOOKUP(6,A2:A11,B2:B11)'
ws['D3'] = '=IFS(A2>5,"büyük",TRUE,"küçük")'
wb.save(src)
zin = zipfile.ZipFile(src)
zout = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == 'xl/worksheets/sheet1.xml':
        x = data.decode()
        # B2: ana paylaşılan formül, B3..B11 bağlı hücreler; önbellek değerleri ekle
        def repl(m):
            r = int(m.group(1)); val = r * 3 * 2
            if r == 2:
                return f'<c r="B2"><f t="shared" ref="B2:B11" si="0">A2*2</f><v>{val}</v></c>'
            return f'<c r="B{r}"><f t="shared" si="0"/><v>{val}</v></c>'
        x = re.sub(r'<c r="B(\d+)"><f>A\d+\*2</f><v></v></c>', repl, x)
        x = re.sub(r'<c r="B(\d+)"><f>A\d+\*2</f><v\s*/></c>', repl, x)
        x = re.sub(r'<c r="B(\d+)"><f>A\d+\*2</f></c>', repl, x)
        x = x.replace('<f>XLOOKUP(', '<f>_xlfn.XLOOKUP(').replace('<f>IFS(', '<f>_xlfn.IFS(')
        data = x.encode()
    zout.writestr(item, data)
zout.close()
print('saved', out)
