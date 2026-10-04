#!/usr/bin/env python3
"""Sınama belgeleri: ../../ewreka-ofis/test-files/ornek.docx (zengin içerik) ve uzun.docx (~44 sayfa)."""
import io, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from PIL import Image, ImageDraw

OUT = os.path.join(os.path.dirname(__file__), '..', '..', '..', 'ewreka-ofis', 'test-files')
os.makedirs(OUT, exist_ok=True)
img = Image.new('RGB', (480, 240), (76, 141, 255)); d = ImageDraw.Draw(img)
d.rectangle([20, 20, 460, 220], outline=(255, 195, 0), width=8); d.ellipse([180, 60, 300, 180], fill=(255, 195, 0))
bio = io.BytesIO(); img.save(bio, 'PNG'); bio.seek(0)

doc = Document(); s = doc.sections[0]
s.page_width, s.page_height = Cm(21), Cm(29.7)
s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.5)
s.header.paragraphs[0].text = 'Ewreka Digital — Üst Bilgi'
s.footer.paragraphs[0].text = 'Gizli · Altbilgi metni'
doc.add_heading('Ewreka Nota Deneme Belgesi', 0)
doc.add_heading('1. Giriş Bölümü', level=1)
p = doc.add_paragraph('Bu paragraf '); p.add_run('kalın').bold = True; p.add_run(', ')
p.add_run('italik').italic = True; p.add_run(' ve '); p.add_run('altı çizili').underline = True
p.add_run(' metin içerir. Türkçe karakterler: ğüşıöç İĞÜŞÖÇ — "Işık ılık süt içti, çiğ börek yedi."')
r = doc.add_paragraph().add_run('Kırmızı Georgia 14 pt'); r.font.color.rgb = RGBColor(0xC0, 0x20, 0x20); r.font.name = 'Georgia'; r.font.size = Pt(14)
doc.add_heading('1.1 Listeler', level=2)
for t in ['Birinci madde', 'İkinci madde', 'Üçüncü madde']: doc.add_paragraph(t, style='List Bullet')
for t in ['Adım bir', 'Adım iki']: doc.add_paragraph(t, style='List Number')
doc.add_heading('1.2 Tablo', level=2)
t = doc.add_table(rows=3, cols=3); t.style = 'Table Grid'
for i, row in enumerate([('Ürün', 'Adet', 'Fiyat'), ('Çay', '2', '₺40,00'), ('Şeker', '1', '₺25,50')]):
    for j, v in enumerate(row): t.cell(i, j).text = v
doc.add_paragraph('Aşağıda bir görsel var:'); doc.add_picture(bio, width=Cm(8))
c = doc.add_paragraph('Ortalanmış paragraf.'); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
doc.add_heading('2. İkinci Sayfa', level=1)
for i in range(6):
    doc.add_paragraph('Lorem ipsum dolor sit amet, consectetur adipiscing elit. Şimdi ğüşıöç harfleriyle uzun bir paragraf yazıyoruz ki satırlar kaysın ve sayfa düzeni sınansın. ' * 2)
doc.save(os.path.join(OUT, 'ornek.docx'))

doc = Document(); doc.add_heading('Uzun Belge', 0)
for i in range(1, 401):
    if i % 25 == 0: doc.add_heading(f'Bölüm {i // 25}', 1)
    doc.add_paragraph(f'{i}. paragraf: Çok sayfalı belge sınaması için Türkçe metin — ğüşıöç İĞÜŞÖÇ. ' * 4)
doc.save(os.path.join(OUT, 'uzun.docx'))
print('ok')
